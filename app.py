"""
Invoice Bridge
--------------
Reconciles internal system records (Voucher Transactions + Payments + Supplier
Masterdata) against the ETA e-invoice portal export.

Run locally:
    pip install -r requirements.txt
    streamlit run app.py
"""

import io
import re
import xml.etree.ElementTree as ET

import pandas as pd
import streamlit as st

# --------------------------------------------------------------------------
# Page config & style
# --------------------------------------------------------------------------
st.set_page_config(page_title="Invoice Bridge", page_icon="🌉", layout="wide")

CUSTOM_CSS = """
<style>
.stApp {
    background-color: #0f1420;
    color: #e8ecf5;
}
section[data-testid="stSidebar"] {
    background-color: #161d2e;
    border-right: 1px solid #232c42;
}
section[data-testid="stSidebar"] * {
    color: #e8ecf5 !important;
}
h1, h2, h3 {
    color: #35d0c0 !important;
    font-weight: 700;
}
p, span, label, li, div {
    color: #e8ecf5;
}
[data-testid="stMetric"] {
    background-color: #161d2e;
    border: 1px solid #232c42;
    border-radius: 10px;
    padding: 14px 16px;
}
[data-testid="stMetricValue"] {
    color: #35d0c0 !important;
    font-weight: 700;
}
[data-testid="stMetricLabel"] {
    color: #9aa4bf !important;
}
div.stButton > button {
    background-color: #f97316;
    color: #0f1420 !important;
    font-weight: 700;
    border: none;
    border-radius: 8px;
}
div.stButton > button p {
    color: #0f1420 !important;
}
div.stButton > button:hover {
    background-color: #fb923c;
    color: #0f1420 !important;
}
.stDownloadButton > button {
    background-color: #35d0c0;
    color: #0f1420 !important;
    font-weight: 700;
    border: none;
    border-radius: 8px;
}
.stDownloadButton > button p {
    color: #0f1420 !important;
}
[data-testid="stFileUploaderDropzone"] {
    background-color: #161d2e;
    border: 1px dashed #35d0c0;
    border-radius: 8px;
}
[data-testid="stFileUploaderDropzone"] * {
    color: #e8ecf5 !important;
}
div[data-baseweb="input"] input {
    background-color: #161d2e;
    color: #e8ecf5 !important;
}
div[data-baseweb="select"] * {
    color: #0f1420 !important;
}
.stDataFrame, [data-testid="stDataFrame"] {
    background-color: #161d2e;
    border: 1px solid #232c42;
    border-radius: 8px;
}
.banner {
    background: linear-gradient(90deg, #161d2e 0%, #1a2338 100%);
    padding: 18px 24px;
    border-radius: 10px;
    border-left: 5px solid #35d0c0;
    margin-bottom: 20px;
}
.banner p {
    color: #9aa4bf !important;
}
.footer-credit {
    text-align: center;
    color: #6b7590;
    font-size: 13px;
    margin-top: 40px;
}
</style>
"""
st.markdown(CUSTOM_CSS, unsafe_allow_html=True)

# --------------------------------------------------------------------------
# Login (placeholder credentials - extend USERS dict to add more accounts)
# --------------------------------------------------------------------------
USERS = {
    "Mahmoud": "1234",
}


def login_screen():
    st.markdown("<h1 style='text-align:center;'>🌉 Invoice Bridge</h1>", unsafe_allow_html=True)
    st.markdown(
        "<p style='text-align:center;color:#8f86ad;'>Supplier ledger &amp; ETA portal reconciliation</p>",
        unsafe_allow_html=True,
    )
    with st.form("login_form"):
        col1, col2, col3 = st.columns([1, 2, 1])
        with col2:
            username = st.text_input("Username")
            password = st.text_input("Password", type="password")
            submitted = st.form_submit_button("Log in", use_container_width=True)
        if submitted:
            if USERS.get(username) == password:
                st.session_state["logged_in"] = True
                st.session_state["username"] = username
                st.rerun()
            else:
                st.error("Invalid username or password.")


if "logged_in" not in st.session_state:
    st.session_state["logged_in"] = False

if not st.session_state["logged_in"]:
    login_screen()
    st.stop()

# --------------------------------------------------------------------------
# Constants / parsing helpers
# --------------------------------------------------------------------------
ACCOUNT_VENDOR_TOTAL = "21020102"
ACCOUNT_VAT = "21030309"
ACCOUNTS_SALES_AMOUNT = {"32000001", "32000002", "32000003", "32000004"}
TARGET_ACCOUNTS = {ACCOUNT_VENDOR_TOTAL, ACCOUNT_VAT} | ACCOUNTS_SALES_AMOUNT

PLACEHOLDER_INVOICE_VALUES = {
    "", "NA", "N/A", "N/ A", ".", "000", "0000000000", "NONE", "-",
}


def clean_str(v):
    if v is None or (isinstance(v, float) and pd.isna(v)) or (isinstance(v, str) and v.strip().lower() == "nan"):
        return ""
    try:
        if pd.isna(v):
            return ""
    except (TypeError, ValueError):
        pass
    return str(v).strip()


def clean_id(v):
    """Normalize ID-like fields (account codes) that pandas may read as floats (e.g. 925.0 -> '925')."""
    s = clean_str(v)
    if s == "":
        return s
    if re.fullmatch(r"-?\d+\.0+", s):
        s = s.split(".")[0]
    return s


def is_placeholder_invoice(v):
    v = clean_str(v).upper().replace(" ", "")
    if v in PLACEHOLDER_INVOICE_VALUES:
        return True
    if v != "" and set(v) == {"0"}:
        return True
    return False


def split_document(doc):
    """Split the 'Document' field into (sales_invoice, vendor_invoice_raw)."""
    doc = clean_str(doc)
    parts = doc.split("/")
    sales_invoice = parts[0] if parts else doc
    vendor_invoice_raw = parts[1] if len(parts) > 1 else ""
    return sales_invoice, vendor_invoice_raw


def read_spreadsheetml_xls(file_bytes):
    """Read old-style SpreadsheetML XML .xls exports (e.g. TINA reports)."""
    ns = {"ss": "urn:schemas-microsoft-com:office:spreadsheet"}
    tree = ET.parse(io.BytesIO(file_bytes))
    root = tree.getroot()
    ws = root.find("ss:Worksheet", ns)
    table = ws.find("ss:Table", ns)
    rows = table.findall("ss:Row", ns)

    data_rows = []
    header = None
    for row in rows:
        vals = []
        for cell in row.findall("ss:Cell", ns):
            data = cell.find("ss:Data", ns)
            vals.append(data.text if data is not None else None)
        if header is None:
            header = vals
        else:
            # pad short rows
            if len(vals) < len(header):
                vals = vals + [None] * (len(header) - len(vals))
            data_rows.append(vals[: len(header)])
    return pd.DataFrame(data_rows, columns=header)


def load_any_excel(uploaded_file):
    """Load .xlsx normally, fall back to SpreadsheetML parser for old .xls exports."""
    name = uploaded_file.name.lower()
    raw = uploaded_file.getvalue()
    if name.endswith(".xls"):
        try:
            return read_spreadsheetml_xls(raw)
        except Exception:
            return pd.read_excel(io.BytesIO(raw))
    return pd.read_excel(io.BytesIO(raw))


# --------------------------------------------------------------------------
# Core processing
# --------------------------------------------------------------------------
def process_voucher_transactions(df):
    """Line-level amounts per (Sales Invoice, Vendor Account) - before re-grouping by vendor invoice no."""
    df = df.copy()
    df["Main account"] = df["Main account"].apply(clean_id)
    df["Vendor account"] = df["Vendor account"].apply(clean_id)
    df["Document"] = df["Document"].apply(clean_str)
    df["Amount in transaction currency"] = pd.to_numeric(
        df["Amount in transaction currency"], errors="coerce"
    ).fillna(0)

    df = df[df["Main account"].isin(TARGET_ACCOUNTS)]
    # drop the aggregate/contra VAT rows that carry no vendor account
    df = df[df["Vendor account"] != ""]

    split = df["Document"].apply(split_document)
    df["Sales Invoice"] = split.apply(lambda t: t[0])
    df["Vendor Invoice Raw"] = split.apply(lambda t: t[1])
    df["Vendor Invoice (Voucher)"] = df["Vendor Invoice Raw"].apply(
        lambda v: "" if is_placeholder_invoice(v) else v
    )

    rows = []
    for (sales_inv, vendor_acc), g in df.groupby(["Sales Invoice", "Vendor account"]):
        vendor_total = abs(g.loc[g["Main account"] == ACCOUNT_VENDOR_TOTAL, "Amount in transaction currency"].sum())
        vat_amount = abs(g.loc[g["Main account"] == ACCOUNT_VAT, "Amount in transaction currency"].sum())
        sales_amount = abs(g.loc[g["Main account"].isin(ACCOUNTS_SALES_AMOUNT), "Amount in transaction currency"].sum())

        voucher_invoice_candidates = [v for v in g["Vendor Invoice (Voucher)"] if v]
        voucher_invoice_no = voucher_invoice_candidates[0] if voucher_invoice_candidates else ""

        rows.append(
            {
                "Sales Invoice": sales_inv,
                "Vendor Account": vendor_acc,
                "Sales Amount": round(sales_amount, 2),
                "VAT Amount": round(vat_amount, 2),
                "Vendor Invoice Total": round(vendor_total, 2),
                "Vendor Invoice No (Voucher)": voucher_invoice_no,
            }
        )

    return pd.DataFrame(rows)


def build_payment_lookup(df):
    df = df.copy()
    df["Vendor account"] = df["Vendor account"].apply(clean_id)
    df["Invoice"] = df["Invoice"].apply(clean_str)
    df["Payment reference"] = df["Payment reference"].apply(clean_str)

    df["Sales Invoice"] = df["Invoice"].apply(lambda v: v.split("/")[0] if v else "")

    lookup = {}
    for _, row in df.iterrows():
        ref = row["Payment reference"]
        if not ref or is_placeholder_invoice(ref):
            continue
        key = (row["Sales Invoice"], row["Vendor account"])
        if key not in lookup:
            lookup[key] = ref
    return lookup


def build_masterdata_lookup(df):
    df = df.copy()
    df["Supplier id"] = df["Supplier id"].apply(clean_id)
    df["Fiscal code"] = df["Fiscal code"].apply(clean_str)
    return dict(zip(df["Supplier id"], df["Fiscal code"]))


def build_stage1_output(voucher_df, payments_df, masterdata_df):
    line_level = process_voucher_transactions(voucher_df)
    payment_lookup = build_payment_lookup(payments_df)
    tax_id_lookup = build_masterdata_lookup(masterdata_df)

    line_level["Vendor Invoice No (Payment)"] = line_level.apply(
        lambda r: payment_lookup.get((r["Sales Invoice"], r["Vendor Account"]), ""), axis=1
    )

    # Effective invoice number used to re-group lines that belong to the same
    # vendor invoice but were split across several of our sales invoices.
    def effective_invoice_no(r):
        if r["Vendor Invoice No (Voucher)"]:
            return r["Vendor Invoice No (Voucher)"]
        if r["Vendor Invoice No (Payment)"]:
            return r["Vendor Invoice No (Payment)"]
        return ""

    line_level["Effective Invoice No"] = line_level.apply(effective_invoice_no, axis=1)

    # Group key: known vendor invoice no -> merge across sales invoices.
    # Unknown vendor invoice no -> keep each (Sales Invoice, Vendor) separate.
    def group_key(r):
        if r["Effective Invoice No"]:
            return (r["Effective Invoice No"], r["Vendor Account"])
        return (f"__NOINV__{r['Sales Invoice']}", r["Vendor Account"])

    line_level["Group Key"] = line_level.apply(group_key, axis=1)

    final_rows = []
    for _, g in line_level.groupby("Group Key"):
        vendor_acc = g["Vendor Account"].iloc[0]
        sales_invoices = "/".join(dict.fromkeys(g["Sales Invoice"]))
        voucher_nos = [v for v in g["Vendor Invoice No (Voucher)"] if v]
        payment_nos = [v for v in g["Vendor Invoice No (Payment)"] if v]

        final_rows.append(
            {
                "Vendor Account": vendor_acc,
                "Tax ID": tax_id_lookup.get(vendor_acc, ""),
                "Sales Invoice": sales_invoices,
                "Vendor Invoice No (Voucher)": voucher_nos[0] if voucher_nos else "",
                "Vendor Invoice No (Payment)": payment_nos[0] if payment_nos else "",
                "COGS": round(g["Sales Amount"].sum(), 2),
                "SUPPLIERS VAT": round(g["VAT Amount"].sum(), 2),
                "Vendor Invoice Total": round(g["Vendor Invoice Total"].sum(), 2),
            }
        )

    return pd.DataFrame(final_rows)


PORTAL_FEE_ADD_COLUMNS = [
    "رسم خدمة",
    "رسم المحليات",
    "رسوم أخرى",
    "ضريبة الدمغة النسبية",
    "ضريبة الدمغة قطعية بمقدار ثابت.1",
    "رسم تنمية الموارد.1",
    "رسم خدمة.1",
    "رسم المحليات.1",
    "رسوم أخرى.1",
]
PORTAL_FEE_SUBTRACT_COLUMNS = ["خصم الفاتورة", "خصم الأصناف"]


def compare_with_eta(stage1_df, eta_df):
    eta = eta_df.copy()
    eta["الرقم الضريبى للبائع"] = eta["الرقم الضريبى للبائع"].apply(clean_id)
    eta["الرقم الداخلى"] = eta["الرقم الداخلى"].apply(clean_id)

    eta_index = {}
    for _, row in eta.iterrows():
        key = (row["الرقم الضريبى للبائع"], row["الرقم الداخلى"])
        eta_index[key] = row

    results = []
    for _, row in stage1_df.iterrows():
        tax_id = row["Tax ID"]
        candidates = [row["Vendor Invoice No (Voucher)"], row["Vendor Invoice No (Payment)"]]
        match = None
        matched_invoice_no = ""
        for cand in candidates:
            if not cand:
                continue
            key = (tax_id, cand)
            if key in eta_index:
                match = eta_index[key]
                matched_invoice_no = cand
                break

        result = row.to_dict()
        if match is not None:
            def _num(v):
                try:
                    if v is None or v == "":
                        return 0.0
                    f = float(v)
                    return 0.0 if pd.isna(f) else f
                except (TypeError, ValueError):
                    return 0.0

            eta_total = _num(match.get("إجمالى الفاتورة"))
            eta_sales = _num(match.get("إجمالى المبيعات"))
            eta_vat = _num(match.get("ضريبة القيمة المضافة"))
            portal_fees = round(
                sum(_num(match.get(c)) for c in PORTAL_FEE_ADD_COLUMNS)
                - sum(_num(match.get(c)) for c in PORTAL_FEE_SUBTRACT_COLUMNS),
                2,
            )

            result["Match Status"] = "Matched"
            result["Vendor Invoice No (Portal Matched)"] = matched_invoice_no
            result["إجمالى المبيعات"] = eta_sales
            result["ضريبة القيمة المضافة"] = eta_vat
            result["ETA Invoice Total"] = eta_total
            result["Portal Fees & Deductions"] = portal_fees
            result["COGS vs إجمالى المبيعات (Difference)"] = round(
                _num(result["COGS"]) - (eta_sales + portal_fees), 2
            )
            result["SUPPLIERS VAT vs ضريبة القيمة المضافة (Difference)"] = round(
                _num(result["SUPPLIERS VAT"]) - eta_vat, 2
            )
            result["Amount Difference"] = round(_num(result["Vendor Invoice Total"]) - eta_total, 2)
        else:
            result["Match Status"] = "Not Found in Portal"
            result["Vendor Invoice No (Portal Matched)"] = ""
            result["إجمالى المبيعات"] = ""
            result["ضريبة القيمة المضافة"] = ""
            result["ETA Invoice Total"] = ""
            result["Portal Fees & Deductions"] = ""
            result["COGS vs إجمالى المبيعات (Difference)"] = ""
            result["SUPPLIERS VAT vs ضريبة القيمة المضافة (Difference)"] = ""
            result["Amount Difference"] = ""
        results.append(result)

    return pd.DataFrame(results)


# --------------------------------------------------------------------------
# UI
# --------------------------------------------------------------------------
st.markdown(
    "<div class='banner'><h1 style='margin:0;'>🌉 Invoice Bridge</h1>"
    "<p style='margin:4px 0 0 0;color:#c9c2e8;'>Supplier ledger reconciliation against the ETA e-invoice portal</p></div>",
    unsafe_allow_html=True,
)

with st.sidebar:
    st.markdown(f"**Logged in as:** {st.session_state.get('username', '')}")
    if st.button("Log out", use_container_width=True):
        st.session_state["logged_in"] = False
        st.rerun()
    st.markdown("---")
    st.markdown("### 1. Upload files")
    masterdata_file = st.file_uploader("Supplier Masterdata (.xls / .xlsx)", type=["xls", "xlsx"])
    voucher_file = st.file_uploader("Voucher Transactions (.xlsx)", type=["xlsx"])
    payments_file = st.file_uploader("Payments (.xlsx)", type=["xlsx"])
    eta_file = st.file_uploader("ETA Portal Export (.xlsx)", type=["xlsx"])
    run_btn = st.button("Run reconciliation", use_container_width=True)

if run_btn:
    if not all([masterdata_file, voucher_file, payments_file, eta_file]):
        st.warning("Please upload all four files before running.")
        st.stop()

    with st.spinner("Reading files..."):
        masterdata_df = load_any_excel(masterdata_file)
        voucher_df = load_any_excel(voucher_file)
        payments_df = load_any_excel(payments_file)
        eta_df = pd.read_excel(io.BytesIO(eta_file.getvalue()), sheet_name="جميع الفواتير")

    with st.spinner("Building system-side reconciliation..."):
        stage1_df = build_stage1_output(voucher_df, payments_df, masterdata_df)

    with st.spinner("Comparing against ETA portal..."):
        final_df = compare_with_eta(stage1_df, eta_df)

    st.session_state["final_df"] = final_df

if "final_df" in st.session_state:
    final_df = st.session_state["final_df"]

    matched = (final_df["Match Status"] == "Matched").sum()
    not_found = (final_df["Match Status"] == "Not Found in Portal").sum()

    c1, c2, c3 = st.columns(3)
    c1.metric("Total Vendor Invoices", len(final_df))
    c2.metric("Matched with Portal", int(matched))
    c3.metric("Not Found in Portal", int(not_found))

    display_cols = [
        "Vendor Account", "Tax ID", "Sales Invoice",
        "Vendor Invoice No (Voucher)", "Vendor Invoice No (Payment)", "Vendor Invoice No (Portal Matched)",
        "COGS", "إجمالى المبيعات", "Portal Fees & Deductions", "COGS vs إجمالى المبيعات (Difference)",
        "SUPPLIERS VAT", "ضريبة القيمة المضافة", "SUPPLIERS VAT vs ضريبة القيمة المضافة (Difference)",
        "Vendor Invoice Total", "ETA Invoice Total", "Amount Difference",
        "Match Status",
    ]
    final_df = final_df[[c for c in display_cols if c in final_df.columns]]

    st.markdown("### Results")
    status_filter = st.selectbox("Filter by status", ["All", "Matched", "Not Found in Portal"])
    display_df = final_df if status_filter == "All" else final_df[final_df["Match Status"] == status_filter]
    st.dataframe(display_df, use_container_width=True, height=500)

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        final_df.to_excel(writer, index=False, sheet_name="Reconciliation")
        ws = writer.sheets["Reconciliation"]

        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        # column groups -> (header color, body tint)
        ID_COLS = {
            "Vendor Account", "Tax ID", "Sales Invoice",
            "Vendor Invoice No (Voucher)", "Vendor Invoice No (Payment)",
            "Vendor Invoice No (Portal Matched)",
        }
        COGS_COLS = {"COGS", "إجمالى المبيعات", "Portal Fees & Deductions", "COGS vs إجمالى المبيعات (Difference)"}
        VAT_COLS = {"SUPPLIERS VAT", "ضريبة القيمة المضافة", "SUPPLIERS VAT vs ضريبة القيمة المضافة (Difference)"}
        TOTAL_COLS = {"Vendor Invoice Total", "ETA Invoice Total", "Amount Difference"}
        STATUS_COLS = {"Match Status"}

        GROUP_COLORS = {
            "id": ("2B3350", "EEF1F8"),
            "cogs": ("1F5F73", "DCEFF3"),
            "vat": ("5B3E82", "EAE1F5"),
            "total": ("2F6B4F", "DFF3E7"),
            "status": ("6B3B22", "FBE7DA"),
        }

        def group_of(col_name):
            if col_name in ID_COLS:
                return "id"
            if col_name in COGS_COLS:
                return "cogs"
            if col_name in VAT_COLS:
                return "vat"
            if col_name in TOTAL_COLS:
                return "total"
            if col_name in STATUS_COLS:
                return "status"
            return "id"

        thin_border = Border(*[Side(style="thin", color="D8DEE9")] * 4)
        columns = list(final_df.columns)

        # header row
        for idx, col_name in enumerate(columns, start=1):
            header_hex, _ = GROUP_COLORS[group_of(col_name)]
            cell = ws.cell(row=1, column=idx)
            cell.fill = PatternFill(start_color=header_hex, end_color=header_hex, fill_type="solid")
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
        ws.row_dimensions[1].height = 30

        matched_fill = PatternFill(start_color="DFF5E9", end_color="DFF5E9", fill_type="solid")
        notfound_fill = PatternFill(start_color="FBE7DA", end_color="FBE7DA", fill_type="solid")
        status_col_idx = columns.index("Match Status") + 1

        numeric_cols = set(columns) - ID_COLS - STATUS_COLS
        numeric_col_idx = {columns.index(c) + 1 for c in numeric_cols}

        for row_idx in range(2, ws.max_row + 1):
            status_val = ws.cell(row=row_idx, column=status_col_idx).value
            row_tint_override = matched_fill if status_val == "Matched" else notfound_fill
            for col_idx, col_name in enumerate(columns, start=1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                if col_name in STATUS_COLS:
                    cell.fill = row_tint_override
                    cell.font = Font(bold=True, color="1F6B45" if status_val == "Matched" else "A34A1F")
                else:
                    _, tint = GROUP_COLORS[group_of(col_name)]
                    cell.fill = PatternFill(start_color=tint, end_color=tint, fill_type="solid")
                if col_idx in numeric_col_idx and isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00"
                    cell.alignment = Alignment(horizontal="right")

        # column widths
        for idx, col_name in enumerate(columns, start=1):
            letter = get_column_letter(idx)
            max_len = max([len(str(col_name))] + [len(str(v)) for v in final_df[col_name].astype(str)])
            ws.column_dimensions[letter].width = min(max(max_len + 3, 12), 34)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    st.download_button(
        "Download Excel report",
        data=output.getvalue(),
        file_name="Invoice_Bridge_Reconciliation.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )

st.markdown("<div class='footer-credit'>Developed by Mahmoud Amin</div>", unsafe_allow_html=True)
